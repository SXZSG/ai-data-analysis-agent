from pathlib import Path

import pandas as pd

from src.utils import dataframe_to_markdown_table, dict_to_markdown_table, ensure_output_dir, format_metric


def export_markdown_report(context: dict, output_dir: Path) -> Path:
    ensure_output_dir(output_dir)
    report_path = output_dir / "data_analysis_report.md"

    lines = [
        "# Data Analysis Report",
        "",
        context.get("ai_summary", ""),
        "",
        "## Agent Workflow Logs",
        "",
    ]

    for log in context.get("logs", []):
        lines.append(f"- {log}")

    lines.extend(
        [
            "",
            "## Dataset Overview",
            "",
            f"- Source: `{context['dataset']['source']}`",
            f"- Rows: {context['dataset']['rows']}",
            f"- Columns: {context['dataset']['columns']}",
            f"- Target: `{context['target']['name']}`",
            f"- Task Type: `{context['task_detection']['task_type']}`",
            "",
            "## Column Profile",
            "",
            dict_to_markdown_table(context["profile"]["column_profiles"]),
            "",
        ]
    )

    if context["profile"].get("numeric_summary"):
        numeric_frame = (
            pd.DataFrame(context["profile"]["numeric_summary"])
            .T.reset_index()
            .rename(columns={"index": "column"})
        )
        lines.extend(
            [
                "## Numeric Summary",
                "",
                dataframe_to_markdown_table(numeric_frame),
                "",
            ]
        )

    lines.extend(
        [
            "## Model Evaluation",
            "",
            _model_results_markdown(context["evaluation"]["results"]),
            "",
            f"Best baseline model: `{context['evaluation']['best_model']['model_name']}`",
            "",
            "## Visualizations",
            "",
        ]
    )

    for name, path in context.get("visualizations", {}).items():
        lines.append(f"- {name}: `{path}`")

    report_path.write_text("\n".join(lines), encoding="utf-8")
    return report_path


def export_word_report(context: dict, output_dir: Path) -> Path | None:
    try:
        from docx import Document
        from docx.shared import Inches
    except Exception:
        return None

    ensure_output_dir(output_dir)
    docx_path = output_dir / "data_analysis_report.docx"

    document = Document()
    document.add_heading("Data Analysis Report", level=1)
    document.add_paragraph(context.get("ai_summary", ""))

    document.add_heading("Dataset Overview", level=2)
    document.add_paragraph(f"Rows: {context['dataset']['rows']}")
    document.add_paragraph(f"Columns: {context['dataset']['columns']}")
    document.add_paragraph(f"Target: {context['target']['name']}")
    document.add_paragraph(f"Task Type: {context['task_detection']['task_type']}")

    document.add_heading("Model Evaluation", level=2)
    for item in context["evaluation"]["results"]:
        document.add_paragraph(item["model_name"], style="List Bullet")
        for metric, value in item["metrics"].items():
            document.add_paragraph(f"{metric}: {format_metric(value)}")

    document.add_heading("Visualizations", level=2)
    for name, path in context.get("visualizations", {}).items():
        document.add_paragraph(name)
        if Path(path).exists():
            document.add_picture(str(path), width=Inches(5.5))

    document.save(docx_path)
    return docx_path


def _model_results_markdown(results: list[dict]) -> str:
    rows = []
    for item in results:
        row = {"model": item["model_name"]}
        row.update({metric: format_metric(value) for metric, value in item["metrics"].items()})
        rows.append(row)
    return dict_to_markdown_table(rows)
